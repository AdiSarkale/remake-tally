"""Tenant-scoped HR payroll, calendar, salary and monthly attendance APIs."""
from __future__ import annotations

import calendar
import io
from datetime import date, datetime, time as time_type, timezone
from decimal import Decimal, ROUND_HALF_UP
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, Query, Response, UploadFile
from fastapi.responses import StreamingResponse
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils.datetime import from_excel
from openpyxl.utils.exceptions import InvalidFileException
from zipfile import BadZipFile
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app import models, schemas
from app.api.deps import require_area
from app.db.session import get_db
from app.services.inventory import log_audit

router = APIRouter(prefix="/hr", tags=["hr", "payroll"])
hr_access = Depends(require_area("hr"))
_CENT = Decimal("0.01")


def _dec(value: Any) -> Decimal:
    return Decimal(str(value if value is not None else 0))


def _money(value: Decimal | float | int) -> Decimal:
    return _dec(value).quantize(_CENT, rounding=ROUND_HALF_UP)


def _ensure_settings(db: Session) -> models.HRPayrollSettings:
    row = db.get(models.HRPayrollSettings, 1)
    if row is None:
        row = models.HRPayrollSettings(id=1)
        db.add(row)
        db.flush()
        db.commit()
        db.refresh(row)
    return row


def _week_offs(settings: models.HRPayrollSettings) -> list[int]:
    try:
        values = [int(token) for token in settings.week_off_days.split(",") if token != ""]
        return sorted(set(day for day in values if 0 <= day <= 6))
    except (AttributeError, ValueError):
        return [5, 6]


def _settings_output(row: models.HRPayrollSettings) -> dict[str, Any]:
    return {
        "id": row.id,
        "salary_count_mode": row.salary_count_mode,
        "monthly_salary_days": row.monthly_salary_days,
        "week_off_days": _week_offs(row),
        "updated_at": row.updated_at,
    }


def _calendar_type(day: date, week_off_days: list[int], overrides: dict[date, str]) -> str:
    return overrides.get(day, "week_off" if day.weekday() in week_off_days else "working_day")


def _salary_count_divisor(
    year: int,
    month: int,
    mode: str,
    monthly_salary_days: int,
    week_off_days: list[int],
    overrides: dict[date, str],
) -> int:
    days_in_month = calendar.monthrange(year, month)[1]
    if mode == "calendar_days":
        return days_in_month
    if mode == "fixed_days":
        return max(1, min(31, monthly_salary_days))
    return max(
        1,
        sum(
            _calendar_type(date(year, month, number), week_off_days, overrides) == "working_day"
            for number in range(1, days_in_month + 1)
        ),
    )


def _salary_components(profile: models.EmployeeSalaryProfile | None) -> dict[str, Decimal]:
    zero = Decimal("0.00")
    if profile is None:
        return {
            "annual_salary": zero, "monthly_gross": zero, "basic": zero, "hra": zero,
            "da": zero, "conveyance": zero, "medical_allowance": zero,
            "special_allowance": zero, "employee_pf": zero, "professional_tax": zero, "tds": zero,
        }
    annual = _money(_dec(profile.annual_salary))
    gross = _money(annual / Decimal("12"))
    basic = _money(gross * _dec(profile.basic_pct) / Decimal("100"))
    hra = _money(basic * _dec(profile.hra_pct) / Decimal("100"))
    da = _money(gross * _dec(profile.da_pct) / Decimal("100"))
    conveyance = _money(_dec(profile.conveyance_monthly))
    medical = _money(_dec(profile.medical_allowance_monthly))
    special = _money(gross - basic - hra - da - conveyance - medical)
    return {
        "annual_salary": annual, "monthly_gross": gross, "basic": basic, "hra": hra, "da": da,
        "conveyance": conveyance, "medical_allowance": medical, "special_allowance": special,
        "employee_pf": _money(basic * _dec(profile.pf_pct) / Decimal("100")),
        "professional_tax": _money(_dec(profile.professional_tax_monthly)),
        "tds": _money(_dec(profile.tds_monthly)),
    }


def _calculate_employee_monthly_payroll(
    employee: Any,
    profile: models.EmployeeSalaryProfile | None,
    attendance_rows: list[Any],
    year: int,
    month: int,
    salary_count_mode: str,
    monthly_salary_days: int,
    week_off_days: list[int],
    calendar_overrides: dict[date, str],
) -> dict[str, Any]:
    components = _salary_components(profile)
    divisor = _salary_count_divisor(
        year, month, salary_count_mode, monthly_salary_days, week_off_days, calendar_overrides
    )
    present_days = half_days = absent_days = paid_leave_days = unpaid_leave_days = Decimal("0")
    unpaid_equivalent_days = Decimal("0")
    overtime_minutes = 0
    for row in attendance_rows:
        # Attendance on an overridden holiday or configured week-off is not a
        # scheduled absence. Still show the recorded attendance in the summary.
        work_date = getattr(row, "work_date", None)
        scheduled_workday = (
            _calendar_type(work_date, week_off_days, calendar_overrides) == "working_day"
            if work_date is not None
            else True
        )
        if row.status == "present":
            present_days += Decimal("1")
        elif row.status == "half_day":
            half_days += Decimal("1")
            if scheduled_workday:
                unpaid_equivalent_days += Decimal("0.5")
        elif row.status == "absent":
            absent_days += Decimal("1")
            if scheduled_workday:
                unpaid_equivalent_days += Decimal("1")
        elif row.status == "paid_leave":
            paid_leave_days += Decimal("1")
        elif row.status == "unpaid_leave":
            unpaid_leave_days += Decimal("1")
            if scheduled_workday:
                unpaid_equivalent_days += Decimal("1")
        # Rejecting or ignoring overtime for absence/leave rows avoids paying
        # overtime that cannot be reconciled to actual attendance.
        if row.status in {"present", "half_day"}:
            overtime_minutes += max(0, int(row.overtime_minutes or 0))

    daily_rate = components["monthly_gross"] / Decimal(divisor)
    absence_deduction = min(
        components["monthly_gross"],
        _money(daily_rate * unpaid_equivalent_days),
    )
    overtime_pay = _money((daily_rate / Decimal("8")) * Decimal(overtime_minutes) / Decimal("60"))
    net_pay = max(
        Decimal("0.00"),
        _money(
            components["monthly_gross"] - absence_deduction - components["employee_pf"]
            - components["professional_tax"] - components["tds"] + overtime_pay
        ),
    )
    return {
        "employee_id": employee.id, "employee_code": employee.emp_code, "employee_name": employee.name,
        "department": employee.department or "", "configured": profile is not None and components["annual_salary"] > 0,
        **components,
        "present_days": _money(present_days), "half_days": _money(half_days),
        "absent_days": _money(absent_days), "paid_leave_days": _money(paid_leave_days),
        "unpaid_leave_days": _money(unpaid_leave_days), "overtime_minutes": overtime_minutes,
        "absence_deduction": absence_deduction, "overtime_pay": overtime_pay, "net_pay": net_pay,
    }


def _validate_period(year: int, month: int) -> None:
    if month < 1 or month > 12 or year < 2000 or year > 2100:
        raise HTTPException(
            status_code=422,
            detail="Payroll period must use a valid month and year from 2000 to 2100",
        )


def _period_data(
    db: Session, year: int, month: int
) -> tuple[date, date, models.HRPayrollSettings, dict[date, str], list[models.Employee]]:
    _validate_period(year, month)
    start = date(year, month, 1)
    end = date(year, month, calendar.monthrange(year, month)[1])
    settings = _ensure_settings(db)
    override_rows = db.query(models.HRCalendarDay).filter(
        models.HRCalendarDay.work_date >= start, models.HRCalendarDay.work_date <= end
    ).all()
    overrides = {row.work_date: row.day_type for row in override_rows}
    employees = db.query(models.Employee).filter(
        models.Employee.active.is_(True)
    ).order_by(models.Employee.emp_code).all()
    return start, end, settings, overrides, employees


def _preview_data(db: Session, year: int, month: int) -> dict[str, Any]:
    start, end, settings, overrides, employees = _period_data(db, year, month)
    week_off_days = _week_offs(settings)
    profiles = {row.employee_id: row for row in db.query(models.EmployeeSalaryProfile).all()}
    attendances = db.query(models.EmployeeAttendance).filter(
        models.EmployeeAttendance.work_date >= start, models.EmployeeAttendance.work_date <= end
    ).all()
    attendance_by_employee: dict[str, list[Any]] = {}
    for row in attendances:
        attendance_by_employee.setdefault(row.employee_id, []).append(row)
    items = [
        _calculate_employee_monthly_payroll(
            employee, profiles.get(employee.id), attendance_by_employee.get(employee.id, []),
            year, month, settings.salary_count_mode, settings.monthly_salary_days,
            week_off_days, overrides,
        )
        for employee in employees
    ]
    salary_count_days = _salary_count_divisor(
        year, month, settings.salary_count_mode, settings.monthly_salary_days, week_off_days, overrides
    )
    return {
        "year": year, "month": month, "salary_count_mode": settings.salary_count_mode,
        "salary_count_days": salary_count_days, "items": items,
    }


def _item_output(item: models.HRPayrollItem) -> dict[str, Any]:
    return {
        "employee_id": item.employee_id, "employee_code": item.employee_code,
        "employee_name": item.employee_name, "department": item.department, "configured": True,
        "annual_salary": float(item.annual_salary), "monthly_gross": float(item.monthly_gross),
        "basic": float(item.basic), "hra": float(item.hra), "da": float(item.da),
        "conveyance": float(item.conveyance), "medical_allowance": float(item.medical_allowance),
        "special_allowance": float(item.special_allowance), "employee_pf": float(item.employee_pf),
        "professional_tax": float(item.professional_tax), "tds": float(item.tds),
        "absence_deduction": float(item.absence_deduction), "overtime_pay": float(item.overtime_pay),
        "net_pay": float(item.net_pay), "present_days": float(item.present_days),
        "half_days": float(item.half_days), "absent_days": float(item.absent_days),
        "paid_leave_days": float(item.paid_leave_days), "unpaid_leave_days": float(item.unpaid_leave_days),
        "overtime_minutes": item.overtime_minutes,
    }


def _run_output(db: Session, run: models.HRPayrollRun, include_items: bool = True) -> dict[str, Any]:
    items = db.query(models.HRPayrollItem).filter(
        models.HRPayrollItem.payroll_run_id == run.id
    ).order_by(models.HRPayrollItem.employee_code).all()
    result: dict[str, Any] = {
        "id": run.id, "year": run.year, "month": run.month, "status": run.status,
        "created_by": run.created_by, "created_at": run.created_at, "finalized_at": run.finalized_at,
        "employee_count": len(items), "net_total": round(sum(float(item.net_pay) for item in items), 2),
    }
    if include_items:
        result["items"] = [_item_output(item) for item in items]
    return result


def _attendance_output(
    employee: models.Employee, work_date: date, row: models.EmployeeAttendance | None,
    calendar_type: str, calendar_description: str,
) -> dict[str, Any]:
    return {
        "id": row.id if row else None, "employee_id": employee.id, "employee_code": employee.emp_code,
        "employee_name": employee.name, "work_date": work_date,
        "status": row.status if row else (calendar_type if calendar_type in {"holiday", "week_off"} else "unmarked"),
        "calendar_type": calendar_type, "calendar_description": calendar_description,
        "punch_in": row.punch_in if row else None, "punch_out": row.punch_out if row else None,
        "overtime_minutes": row.overtime_minutes if row else 0, "note": row.note if row else "",
    }



_ATTENDANCE_EXCEL_HEADERS = [
    "Employee Code", "Work Date", "Status", "Overtime Minutes", "Punch In", "Punch Out", "Note",
]
_ATTENDANCE_EXCEL_STATUSES = {"present", "half_day", "absent", "paid_leave", "unpaid_leave", "unmarked"}
_MAX_ATTENDANCE_IMPORT_BYTES = 5 * 1024 * 1024
_MAX_ATTENDANCE_IMPORT_ROWS = 10000


def _attendance_template_workbook(year: int, month: int) -> io.BytesIO:
    """Build a blank, validated workbook. Example data is kept off the import sheet."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Attendance"
    sheet.append(_ATTENDANCE_EXCEL_HEADERS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:G1"
    widths = [20, 16, 20, 20, 14, 14, 36]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + column)].width = width
        cell = sheet.cell(row=1, column=column)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244062")
    for row_number in range(2, 1002):
        sheet.cell(row=row_number, column=2).number_format = "yyyy-mm-dd"
        sheet.cell(row=row_number, column=4).number_format = "0"
        sheet.cell(row=row_number, column=5).number_format = "hh:mm"
        sheet.cell(row=row_number, column=6).number_format = "hh:mm"
    status_validation = DataValidation(
        type="list",
        formula1='"present,half_day,absent,paid_leave,unpaid_leave,unmarked"',
        allow_blank=False,
    )
    status_validation.errorTitle = "Invalid attendance status"
    status_validation.error = "Choose a status from the dropdown list."
    status_validation.promptTitle = "Attendance status"
    status_validation.prompt = "Choose one of the supported lowercase status values."
    sheet.add_data_validation(status_validation)
    status_validation.add("C2:C1001")

    instructions = workbook.create_sheet("Instructions")
    lines = [
        ["MiniTally Attendance Import Template"],
        [f"Selected import period: {calendar.month_name[month]} {year}"],
        ["Only the Attendance worksheet is imported. The Instructions worksheet is ignored."],
        ["Keep the seven column headers on row 1 of Attendance exactly as provided."],
        ["Use one row per employee per work date. Do not upload duplicate employee/date pairs."],
        ["Employee Code: exact code of an active employee, for example EMP-001."],
        ["Work Date: an Excel date or YYYY-MM-DD; every date must belong to the selected period."],
        ["Status: present, half_day, absent, paid_leave, unpaid_leave, or unmarked."],
        ["Overtime Minutes: whole number from 0 to 1440; leave blank or enter 0 if none."],
        ["Punch In / Punch Out: optional Excel times (HH:MM). Punch Out must be after Punch In."],
        ["Note: optional text describing the attendance entry."],
        ["Use unmarked to clear any existing attendance entry for that employee/date."],
        ["The import validates every row first. If any row is invalid, nothing is imported."],
        [""],
        ["Example (documentation only; this row is NOT imported):"],
        _ATTENDANCE_EXCEL_HEADERS,
        ["EMP-001", date(year, month, 1), "present", 0, time_type(8, 30), time_type(17, 0), "Example only"],
    ]
    for row in lines:
        instructions.append(row)
    instructions["A1"].font = Font(bold=True, size=15, color="244062")
    instructions.column_dimensions["A"].width = 110
    for cell in instructions[16]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244062")
    instructions["B17"].number_format = "yyyy-mm-dd"
    instructions["E17"].number_format = "hh:mm"
    instructions["F17"].number_format = "hh:mm"

    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    buffer.seek(0)
    return buffer


def _parse_attendance_excel_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        raw = value.strip()
        for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
            try:
                return datetime.strptime(raw, fmt).date()
            except ValueError:
                continue
    raise ValueError("Work Date must be an Excel date or YYYY-MM-DD")


def _parse_attendance_excel_time(value: Any) -> time_type | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, datetime):
        return value.time().replace(tzinfo=None)
    if isinstance(value, time_type):
        return value.replace(tzinfo=None)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        converted = from_excel(value)
        if isinstance(converted, datetime):
            return converted.time().replace(tzinfo=None)
        if isinstance(converted, time_type):
            return converted.replace(tzinfo=None)
    if isinstance(value, str):
        raw = value.strip()
        try:
            return time_type.fromisoformat(raw)
        except ValueError:
            for fmt in ("%H:%M", "%H:%M:%S"):
                try:
                    return datetime.strptime(raw, fmt).time()
                except ValueError:
                    continue
    raise ValueError("time must be in HH:MM or HH:MM:SS format")


def _parse_attendance_workbook(
    content: bytes,
    year: int,
    month: int,
    employees_by_code: dict[str, Any],
) -> list[dict[str, Any]]:
    """Validate the entire workbook before the caller makes any database changes."""
    _validate_period(year, month)
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except (InvalidFileException, BadZipFile, OSError, ValueError) as exc:
        raise ValueError("The file is not a readable Excel .xlsx workbook. Download the MiniTally template and try again.") from exc
    try:
        if "Attendance" not in workbook.sheetnames:
            raise ValueError('Workbook must contain a worksheet named "Attendance".')
        sheet = workbook["Attendance"]
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        headers = [str(value).strip() if value is not None else "" for value in (header_row or ())]
        while headers and not headers[-1]:
            headers.pop()
        if headers != _ATTENDANCE_EXCEL_HEADERS:
            expected = ", ".join(_ATTENDANCE_EXCEL_HEADERS)
            raise ValueError(f"Attendance worksheet headers must match the template exactly: {expected}.")

        parsed: list[dict[str, Any]] = []
        errors: list[str] = []
        seen: set[tuple[str, date]] = set()
        nonempty_count = 0
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            cells = list(values[:len(_ATTENDANCE_EXCEL_HEADERS)])
            cells.extend([None] * (len(_ATTENDANCE_EXCEL_HEADERS) - len(cells)))
            if not any(value is not None and (not isinstance(value, str) or value.strip()) for value in cells):
                continue
            nonempty_count += 1
            if nonempty_count > _MAX_ATTENDANCE_IMPORT_ROWS:
                errors.append(f"More than {_MAX_ATTENDANCE_IMPORT_ROWS} attendance rows were supplied.")
                break
            row_error_start = len(errors)
            raw_code, raw_date, raw_status, raw_overtime, raw_punch_in, raw_punch_out, raw_note = cells
            code = str(raw_code).strip().upper() if raw_code is not None else ""
            employee = employees_by_code.get(code)
            if not code:
                errors.append(f"Row {row_number}: Employee Code is required.")
            elif employee is None:
                errors.append(f"Row {row_number}: employee code {code!r} was not found or is inactive.")

            work_date = None
            try:
                work_date = _parse_attendance_excel_date(raw_date)
                if work_date.year != year or work_date.month != month:
                    errors.append(f"Row {row_number}: Work Date must be in {year}-{month:02d}.")
            except ValueError as exc:
                errors.append(f"Row {row_number}: {exc}.")

            status = str(raw_status).strip().lower().replace("-", "_").replace(" ", "_") if raw_status is not None else ""
            if status not in _ATTENDANCE_EXCEL_STATUSES:
                errors.append(f"Row {row_number}: Status must be one of {', '.join(sorted(_ATTENDANCE_EXCEL_STATUSES))}.")

            overtime_minutes = 0
            if raw_overtime is not None and str(raw_overtime).strip():
                try:
                    overtime_number = float(raw_overtime)
                    if isinstance(raw_overtime, bool) or not overtime_number.is_integer() or not 0 <= overtime_number <= 1440:
                        raise ValueError
                    overtime_minutes = int(overtime_number)
                except (TypeError, ValueError):
                    errors.append(f"Row {row_number}: Overtime Minutes must be a whole number from 0 to 1440.")

            punch_in = punch_out = None
            try:
                parsed_time = _parse_attendance_excel_time(raw_punch_in)
                if parsed_time is not None and work_date is not None:
                    punch_in = datetime.combine(work_date, parsed_time).replace(tzinfo=timezone.utc)
            except ValueError as exc:
                errors.append(f"Row {row_number}: Punch In {exc}.")
            try:
                parsed_time = _parse_attendance_excel_time(raw_punch_out)
                if parsed_time is not None and work_date is not None:
                    punch_out = datetime.combine(work_date, parsed_time).replace(tzinfo=timezone.utc)
            except ValueError as exc:
                errors.append(f"Row {row_number}: Punch Out {exc}.")
            if punch_in is not None and punch_out is not None and punch_out <= punch_in:
                errors.append(f"Row {row_number}: Punch Out must be after Punch In.")

            note = str(raw_note).strip() if raw_note is not None else ""
            if len(note) > 4000:
                errors.append(f"Row {row_number}: Note must be 4000 characters or fewer.")

            if work_date is not None and code:
                key = (code, work_date)
                if key in seen:
                    errors.append(f"Row {row_number}: duplicate Employee Code and Work Date pair ({code}, {work_date}).")
                seen.add(key)

            if len(errors) == row_error_start and employee is not None and work_date is not None:
                parsed.append({
                    "employee_id": employee.id,
                    "employee_code": code,
                    "work_date": work_date,
                    "status": status,
                    "overtime_minutes": overtime_minutes,
                    "punch_in": punch_in,
                    "punch_out": punch_out,
                    "note": note,
                })

        if nonempty_count == 0:
            errors.append("Attendance worksheet has no attendance rows.")
        if errors:
            shown = errors[:30]
            suffix = f"\n…and {len(errors) - len(shown)} additional issue(s)." if len(errors) > len(shown) else ""
            raise ValueError("Attendance import validation failed:\n" + "\n".join(shown) + suffix)
        return parsed
    finally:
        workbook.close()


@router.get("/payroll/settings", response_model=schemas.PayrollSettingsOut)
def get_payroll_settings(db: Session = Depends(get_db), user: models.User = hr_access):
    return _settings_output(_ensure_settings(db))


@router.put("/payroll/settings", response_model=schemas.PayrollSettingsOut)
def update_payroll_settings(
    payload: schemas.PayrollSettingsIn, db: Session = Depends(get_db), user: models.User = hr_access
):
    row = _ensure_settings(db)
    row.salary_count_mode = payload.salary_count_mode
    row.monthly_salary_days = payload.monthly_salary_days
    row.week_off_days = ",".join(str(day) for day in payload.week_off_days)
    log_audit(db, user.username, "UPDATE", "hr_payroll_settings", payload.salary_count_mode)
    db.commit()
    db.refresh(row)
    return _settings_output(row)


@router.get("/calendar", response_model=list[schemas.CalendarDayOut])
def get_hr_calendar(
    year: int = Query(ge=2000, le=2100), month: int = Query(ge=1, le=12),
    db: Session = Depends(get_db), user: models.User = hr_access,
):
    start, end, settings, overrides, _ = _period_data(db, year, month)
    rows = {
        row.work_date: row for row in db.query(models.HRCalendarDay).filter(
            models.HRCalendarDay.work_date >= start, models.HRCalendarDay.work_date <= end
        ).all()
    }
    week_off_days = _week_offs(settings)
    return [{
        "work_date": (day := date(year, month, number)),
        "day_type": _calendar_type(day, week_off_days, overrides),
        "description": rows[day].description if day in rows else "",
        "is_override": day in rows,
    } for number in range(1, end.day + 1)]


@router.put("/calendar/{work_date}", response_model=schemas.CalendarDayOut)
def update_hr_calendar_day(
    work_date: date, payload: schemas.CalendarDayIn,
    db: Session = Depends(get_db), user: models.User = hr_access,
):
    _ensure_settings(db)
    row = db.query(models.HRCalendarDay).filter_by(work_date=work_date).first()
    if row is None:
        row = models.HRCalendarDay(work_date=work_date, day_type=payload.day_type, description=payload.description.strip())
        db.add(row)
    else:
        row.day_type = payload.day_type
        row.description = payload.description.strip()
    log_audit(db, user.username, "UPSERT", "hr_calendar_day", f"{work_date}:{payload.day_type}")
    db.commit()
    db.refresh(row)
    return {"work_date": row.work_date, "day_type": row.day_type, "description": row.description, "is_override": True}


@router.delete("/calendar/{work_date}", status_code=204)
def delete_hr_calendar_override(
    work_date: date, db: Session = Depends(get_db), user: models.User = hr_access
):
    row = db.query(models.HRCalendarDay).filter_by(work_date=work_date).first()
    if row is not None:
        db.delete(row)
        log_audit(db, user.username, "DELETE", "hr_calendar_day", str(work_date))
        db.commit()
    return Response(status_code=204)


def _salary_output(employee: models.Employee, profile: models.EmployeeSalaryProfile | None) -> dict[str, Any]:
    components = _salary_components(profile)
    return {
        "employee_id": employee.id, "employee_code": employee.emp_code, "employee_name": employee.name,
        "department": employee.department or "", "active": employee.active, "configured": profile is not None,
        "annual_salary": float(components["annual_salary"]),
        "basic_pct": float(profile.basic_pct) if profile else 50.0,
        "hra_pct": float(profile.hra_pct) if profile else 40.0,
        "da_pct": float(profile.da_pct) if profile else 0.0,
        "conveyance_monthly": float(profile.conveyance_monthly) if profile else 0.0,
        "medical_allowance_monthly": float(profile.medical_allowance_monthly) if profile else 0.0,
        "pf_pct": float(profile.pf_pct) if profile else 12.0,
        "professional_tax_monthly": float(profile.professional_tax_monthly) if profile else 200.0,
        "tds_monthly": float(profile.tds_monthly) if profile else 0.0,
        "monthly_gross": float(components["monthly_gross"]), "basic": float(components["basic"]),
        "hra": float(components["hra"]), "da": float(components["da"]),
        "conveyance": float(components["conveyance"]), "medical_allowance": float(components["medical_allowance"]),
        "special_allowance": float(components["special_allowance"]), "employee_pf": float(components["employee_pf"]),
        "professional_tax": float(components["professional_tax"]), "tds": float(components["tds"]),
        "updated_at": profile.updated_at if profile else None,
    }


@router.get("/salaries", response_model=list[schemas.EmployeeSalaryOut])
def list_employee_salaries(db: Session = Depends(get_db), user: models.User = hr_access):
    # Salary and attendance workflows operate on active employees. Inactive
    # employees remain available in the employee history UI, but shouldn't
    # reappear in the active Payroll workspace after a demo reset/deactivation.
    employees = db.query(models.Employee).filter(
        models.Employee.active.is_(True)
    ).order_by(models.Employee.emp_code).all()
    profiles = {row.employee_id: row for row in db.query(models.EmployeeSalaryProfile).all()}
    return [_salary_output(employee, profiles.get(employee.id)) for employee in employees]


@router.put("/employees/{employee_id}/salary", response_model=schemas.EmployeeSalaryOut)
def upsert_employee_salary(
    employee_id: str, payload: schemas.EmployeeSalaryIn,
    db: Session = Depends(get_db), user: models.User = hr_access,
):
    employee = db.get(models.Employee, employee_id)
    if employee is None:
        raise HTTPException(status_code=404, detail="Employee not found")
    annual = _dec(payload.annual_salary)
    gross = annual / Decimal("12")
    basic = gross * _dec(payload.basic_pct) / Decimal("100")
    hra = basic * _dec(payload.hra_pct) / Decimal("100")
    da = gross * _dec(payload.da_pct) / Decimal("100")
    fixed = _dec(payload.conveyance_monthly) + _dec(payload.medical_allowance_monthly)
    if basic + hra + da + fixed > gross:
        raise HTTPException(status_code=422, detail="Salary components exceed monthly gross salary; reduce percentages or fixed allowances")
    profile = db.get(models.EmployeeSalaryProfile, employee_id)
    if profile is None:
        profile = models.EmployeeSalaryProfile(employee_id=employee_id, annual_salary=payload.annual_salary)
        db.add(profile)
    for field in (
        "annual_salary", "basic_pct", "hra_pct", "da_pct", "conveyance_monthly",
        "medical_allowance_monthly", "pf_pct", "professional_tax_monthly", "tds_monthly",
    ):
        setattr(profile, field, getattr(payload, field))
    log_audit(db, user.username, "UPSERT", "employee_salary_profile", employee.emp_code)
    db.commit()
    db.refresh(profile)
    return _salary_output(employee, profile)


@router.get("/attendance/monthly", response_model=list[schemas.EmployeeAttendanceOut])
def get_monthly_attendance(
    employee_id: str, year: int = Query(ge=2000, le=2100), month: int = Query(ge=1, le=12),
    db: Session = Depends(get_db), user: models.User = hr_access,
):
    start, end, settings, overrides, _ = _period_data(db, year, month)
    employee = db.get(models.Employee, employee_id)
    if employee is None:
        raise HTTPException(status_code=404, detail="Employee not found")
    records = db.query(models.EmployeeAttendance).filter(
        models.EmployeeAttendance.employee_id == employee_id,
        models.EmployeeAttendance.work_date >= start, models.EmployeeAttendance.work_date <= end,
    ).all()
    by_date = {record.work_date: record for record in records}
    calendar_rows = db.query(models.HRCalendarDay).filter(
        models.HRCalendarDay.work_date >= start, models.HRCalendarDay.work_date <= end
    ).all()
    descriptions = {row.work_date: row.description for row in calendar_rows}
    week_off_days = _week_offs(settings)
    return [
        _attendance_output(employee, day, by_date.get(day), _calendar_type(day, week_off_days, overrides), descriptions.get(day, ""))
        for day in (date(year, month, number) for number in range(1, end.day + 1))
    ]



@router.get("/attendance/template")
def download_attendance_template(
    year: int = Query(ge=2000, le=2100),
    month: int = Query(ge=1, le=12),
    db: Session = Depends(get_db),
    user: models.User = hr_access,
):
    _validate_period(year, month)
    workbook = _attendance_template_workbook(year, month)
    filename = f"minitally-attendance-template-{year}-{month:02d}.xlsx"
    return StreamingResponse(
        workbook,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.post("/attendance/import")
async def import_attendance_from_excel(
    year: int = Query(ge=2000, le=2100),
    month: int = Query(ge=1, le=12),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: models.User = hr_access,
):
    _validate_period(year, month)
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=415, detail="Upload an .xlsx Excel workbook using the MiniTally attendance template.")
    content = await file.read(_MAX_ATTENDANCE_IMPORT_BYTES + 1)
    await file.close()
    if len(content) > _MAX_ATTENDANCE_IMPORT_BYTES:
        raise HTTPException(status_code=413, detail="Attendance workbook must be 5 MB or smaller.")

    employees_by_code = {
        row.emp_code.strip().upper(): row
        for row in db.query(models.Employee).filter(models.Employee.active.is_(True)).all()
    }
    try:
        rows = _parse_attendance_workbook(content, year, month, employees_by_code)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc

    created = updated = cleared = 0
    try:
        for item in rows:
            existing = db.query(models.EmployeeAttendance).filter_by(
                employee_id=item["employee_id"], work_date=item["work_date"]
            ).first()
            if item["status"] == "unmarked":
                if existing is not None:
                    db.delete(existing)
                    cleared += 1
                continue
            if existing is None:
                existing = models.EmployeeAttendance(
                    employee_id=item["employee_id"], work_date=item["work_date"],
                    status=item["status"], overtime_minutes=item["overtime_minutes"],
                    punch_in=item["punch_in"], punch_out=item["punch_out"], note=item["note"],
                )
                db.add(existing)
                created += 1
            else:
                existing.status = item["status"]
                existing.overtime_minutes = item["overtime_minutes"]
                existing.punch_in = item["punch_in"]
                existing.punch_out = item["punch_out"]
                existing.note = item["note"]
                updated += 1
        log_audit(
            db, user.username, "IMPORT", "employee_attendance",
            f"{year}-{month:02d}: rows={len(rows)}, created={created}, updated={updated}, cleared={cleared}",
        )
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="Attendance changed concurrently. Refresh the month and import again.") from exc

    return {
        "year": year, "month": month, "rows_processed": len(rows),
        "created": created, "updated": updated, "cleared": cleared,
    }


@router.put("/attendance", response_model=schemas.EmployeeAttendanceOut)
def upsert_daily_attendance(
    payload: schemas.EmployeeAttendanceIn, db: Session = Depends(get_db), user: models.User = hr_access
):
    employee = db.get(models.Employee, payload.employee_id)
    if employee is None:
        raise HTTPException(status_code=404, detail="Employee not found")
    punch_in, punch_out = payload.punch_in, payload.punch_out
    if punch_in is not None and punch_in.tzinfo is None:
        punch_in = punch_in.replace(tzinfo=timezone.utc)
    if punch_out is not None and punch_out.tzinfo is None:
        punch_out = punch_out.replace(tzinfo=timezone.utc)
    if punch_in and punch_out and punch_out <= punch_in:
        raise HTTPException(status_code=422, detail="Punch-out must be after punch-in")
    row = db.query(models.EmployeeAttendance).filter_by(employee_id=payload.employee_id, work_date=payload.work_date).first()
    if payload.status == "unmarked":
        if row is not None:
            db.delete(row)
    elif row is None:
        row = models.EmployeeAttendance(
            employee_id=payload.employee_id, work_date=payload.work_date, status=payload.status,
            punch_in=punch_in, punch_out=punch_out, overtime_minutes=payload.overtime_minutes, note=payload.note.strip(),
        )
        db.add(row)
    else:
        row.status = payload.status
        row.punch_in = punch_in
        row.punch_out = punch_out
        row.overtime_minutes = payload.overtime_minutes
        row.note = payload.note.strip()
    log_audit(db, user.username, "UPSERT", "employee_attendance", f"{employee.emp_code}:{payload.work_date}:{payload.status}")
    db.commit()
    row = db.query(models.EmployeeAttendance).filter_by(employee_id=payload.employee_id, work_date=payload.work_date).first()
    settings = _ensure_settings(db)
    override = db.query(models.HRCalendarDay).filter_by(work_date=payload.work_date).first()
    calendar_type = _calendar_type(payload.work_date, _week_offs(settings), {payload.work_date: override.day_type} if override else {})
    return _attendance_output(employee, payload.work_date, row, calendar_type, override.description if override else "")


@router.delete("/attendance/{employee_id}/{work_date}", status_code=204)
def clear_daily_attendance(
    employee_id: str, work_date: date, db: Session = Depends(get_db), user: models.User = hr_access
):
    employee = db.get(models.Employee, employee_id)
    if employee is None:
        raise HTTPException(status_code=404, detail="Employee not found")
    row = db.query(models.EmployeeAttendance).filter_by(employee_id=employee_id, work_date=work_date).first()
    if row is not None:
        db.delete(row)
        log_audit(db, user.username, "DELETE", "employee_attendance", f"{employee.emp_code}:{work_date}")
        db.commit()
    return Response(status_code=204)


@router.get("/payroll/preview", response_model=schemas.PayrollPreviewOut)
def preview_monthly_payroll(
    year: int = Query(ge=2000, le=2100), month: int = Query(ge=1, le=12),
    db: Session = Depends(get_db), user: models.User = hr_access,
):
    return _preview_data(db, year, month)


@router.get("/payroll/runs", response_model=list[schemas.PayrollRunSummaryOut])
def list_payroll_runs(db: Session = Depends(get_db), user: models.User = hr_access):
    runs = db.query(models.HRPayrollRun).order_by(models.HRPayrollRun.year.desc(), models.HRPayrollRun.month.desc()).all()
    return [_run_output(db, run, include_items=False) for run in runs]


@router.get("/payroll/runs/{year}/{month}", response_model=schemas.PayrollRunOut)
def get_payroll_run(
    year: int, month: int, db: Session = Depends(get_db), user: models.User = hr_access
):
    _validate_period(year, month)
    run = db.query(models.HRPayrollRun).filter_by(year=year, month=month).first()
    if run is None:
        raise HTTPException(status_code=404, detail="Finalized payroll run not found")
    return _run_output(db, run)


@router.post("/payroll/runs/{year}/{month}/finalize", response_model=schemas.PayrollRunOut)
def finalize_monthly_payroll(
    year: int, month: int, db: Session = Depends(get_db), user: models.User = hr_access
):
    _validate_period(year, month)
    existing = db.query(models.HRPayrollRun).filter_by(year=year, month=month).first()
    if existing is not None and existing.status == "finalized":
        return _run_output(db, existing)
    preview = _preview_data(db, year, month)
    configured_items = [item for item in preview["items"] if item["configured"] and item["annual_salary"] > 0]
    if not configured_items:
        raise HTTPException(status_code=409, detail="Configure annual salary for at least one active employee before finalizing payroll")
    now = datetime.now(timezone.utc)
    run = existing or models.HRPayrollRun(year=year, month=month, status="finalized", created_by=user.username, finalized_at=now)
    run.status, run.finalized_at = "finalized", now
    if existing is None:
        db.add(run)
    try:
        db.flush()
        if existing is not None:
            db.query(models.HRPayrollItem).filter_by(payroll_run_id=run.id).delete(synchronize_session=False)
        for item in configured_items:
            db.add(models.HRPayrollItem(
                payroll_run_id=run.id, employee_id=item["employee_id"],
                employee_code=item["employee_code"], employee_name=item["employee_name"],
                department=item["department"], annual_salary=item["annual_salary"],
                monthly_gross=item["monthly_gross"], basic=item["basic"], hra=item["hra"], da=item["da"],
                conveyance=item["conveyance"], medical_allowance=item["medical_allowance"],
                special_allowance=item["special_allowance"], employee_pf=item["employee_pf"],
                professional_tax=item["professional_tax"], tds=item["tds"],
                absence_deduction=item["absence_deduction"], overtime_pay=item["overtime_pay"],
                net_pay=item["net_pay"], present_days=item["present_days"], half_days=item["half_days"],
                absent_days=item["absent_days"], paid_leave_days=item["paid_leave_days"],
                unpaid_leave_days=item["unpaid_leave_days"], overtime_minutes=item["overtime_minutes"],
            ))
        log_audit(db, user.username, "FINALIZE", "hr_payroll_run", f"{year}-{month:02d}")
        db.commit()
        db.refresh(run)
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="Payroll was finalized concurrently for this month; refresh the payroll run list") from exc
    return _run_output(db, run)


@router.get("/payroll/slips/{employee_id}/{year}/{month}")
def download_salary_slip(
    employee_id: str, year: int, month: int, db: Session = Depends(get_db), user: models.User = hr_access
):
    _validate_period(year, month)
    run = db.query(models.HRPayrollRun).filter_by(year=year, month=month, status="finalized").first()
    if run is None:
        raise HTTPException(status_code=404, detail="Finalize payroll for this month before downloading salary slips")
    item = db.query(models.HRPayrollItem).filter_by(payroll_run_id=run.id, employee_id=employee_id).first()
    if item is None:
        raise HTTPException(status_code=404, detail="Salary slip not found for this employee and period")

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    pdf.setTitle(f"Salary Slip {item.employee_code} {year}-{month:02d}")
    pdf.setFont("Helvetica-Bold", 17)
    pdf.drawString(48, 800, "MINITALLY — SALARY SLIP")
    lines = [
        ("Employee", item.employee_name), ("Employee Code", item.employee_code),
        ("Department", item.department or "—"), ("Salary Period", f"{calendar.month_name[month]} {year}"),
        ("Annual Salary", f"INR {float(item.annual_salary):,.2f}"),
        ("Monthly Gross", f"INR {float(item.monthly_gross):,.2f}"),
        ("Basic", f"INR {float(item.basic):,.2f}"), ("HRA", f"INR {float(item.hra):,.2f}"),
        ("DA", f"INR {float(item.da):,.2f}"), ("Conveyance", f"INR {float(item.conveyance):,.2f}"),
        ("Medical Allowance", f"INR {float(item.medical_allowance):,.2f}"),
        ("Special Allowance", f"INR {float(item.special_allowance):,.2f}"),
        ("Absence / Unpaid Leave", f"- INR {float(item.absence_deduction):,.2f}"),
        ("Employee PF", f"- INR {float(item.employee_pf):,.2f}"),
        ("Professional Tax", f"- INR {float(item.professional_tax):,.2f}"),
        ("TDS", f"- INR {float(item.tds):,.2f}"),
        ("Overtime Pay", f"+ INR {float(item.overtime_pay):,.2f}"),
        ("Present / Half / Absent days", f"{float(item.present_days):g} / {float(item.half_days):g} / {float(item.absent_days):g}"),
        ("Paid / Unpaid Leave", f"{float(item.paid_leave_days):g} / {float(item.unpaid_leave_days):g}"),
        ("Overtime", f"{item.overtime_minutes} minutes"),
        ("NET PAYABLE", f"INR {float(item.net_pay):,.2f}"),
    ]
    y = 765
    for label, value in lines:
        if label == "NET PAYABLE":
            y -= 6
            pdf.line(48, y + 12, 545, y + 12)
            pdf.setFont("Helvetica-Bold", 12)
        else:
            pdf.setFont("Helvetica", 10)
        pdf.drawString(55, y, label)
        pdf.drawRightString(535, y, value)
        y -= 26
    pdf.setFont("Helvetica-Oblique", 8)
    pdf.drawString(48, 35, "System-generated salary slip. Verify payroll values with your HR/payroll administrator.")
    pdf.save()
    buffer.seek(0)
    log_audit(
        db,
        user.username,
        "DOWNLOAD",
        "hr_salary_slip",
        f"{year}-{month:02d}:{item.employee_code}",
    )
    db.commit()
    filename = f"salary-slip-{item.employee_code}-{year}-{month:02d}.pdf"
    return StreamingResponse(buffer, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{filename}"'})
