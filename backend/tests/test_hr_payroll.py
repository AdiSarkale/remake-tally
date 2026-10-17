"""Tests for salary composition, calendar rules and payroll calculations."""
from datetime import date, time as time_type
from io import BytesIO
from decimal import Decimal
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from app.api.routes.hr_payroll import (
    _attendance_template_workbook,
    _calculate_employee_monthly_payroll,
    _calendar_type,
    _parse_attendance_workbook,
    _salary_count_divisor,
)
from app.modules.hr.payroll_schemas import PayrollSettingsIn


def _employee():
    return SimpleNamespace(id="emp-1", emp_code="EMP-001", name="Test Employee", department="Production")


def _profile():
    return SimpleNamespace(
        annual_salary=600000, basic_pct=50, hra_pct=40, da_pct=0,
        conveyance_monthly=1000, medical_allowance_monthly=1000,
        pf_pct=12, professional_tax_monthly=200, tds_monthly=500,
    )


def test_salary_breakup_sums_to_monthly_gross_and_computes_deductions():
    from app.api.routes.hr_payroll import _salary_components

    result = _salary_components(_profile())
    assert result["annual_salary"] == 600000
    assert result["monthly_gross"] == 50000
    assert result["basic"] == 25000
    assert result["hra"] == 10000
    assert result["da"] == 0
    assert result["conveyance"] == 1000
    assert result["medical_allowance"] == 1000
    assert result["special_allowance"] == 13000
    assert sum(result[key] for key in ("basic", "hra", "da", "conveyance", "medical_allowance", "special_allowance")) == result["monthly_gross"]
    assert result["employee_pf"] == 3000
    assert result["professional_tax"] == 200
    assert result["tds"] == 500


def test_absence_half_day_leave_and_overtime_affect_net_pay():
    attendance = [
        SimpleNamespace(status="absent", overtime_minutes=0),
        SimpleNamespace(status="half_day", overtime_minutes=0),
        SimpleNamespace(status="unpaid_leave", overtime_minutes=0),
        SimpleNamespace(status="paid_leave", overtime_minutes=0),
        SimpleNamespace(status="present", overtime_minutes=120),
    ]
    result = _calculate_employee_monthly_payroll(
        _employee(), _profile(), attendance, 2026, 10, "fixed_days", 30, [5, 6], {}
    )
    assert result["absent_days"] == 1
    assert result["half_days"] == 1
    assert result["paid_leave_days"] == 1
    assert result["unpaid_leave_days"] == 1
    assert result["overtime_minutes"] == 120
    assert float(result["absence_deduction"]) == 4166.67
    assert float(result["overtime_pay"]) == 416.67
    assert float(result["net_pay"]) == 42550.00


def test_calendar_overrides_take_precedence_over_week_off_defaults():
    overrides = {date(2026, 10, 3): "working_day", date(2026, 10, 5): "holiday"}
    assert _calendar_type(date(2026, 10, 3), [5, 6], overrides) == "working_day"
    assert _calendar_type(date(2026, 10, 4), [5, 6], overrides) == "week_off"
    assert _calendar_type(date(2026, 10, 5), [5, 6], overrides) == "holiday"


def test_salary_count_divisor_supports_fixed_calendar_and_working_days():
    overrides = {date(2026, 10, 3): "working_day", date(2026, 10, 5): "holiday"}
    assert _salary_count_divisor(2026, 10, "calendar_days", 26, [5, 6], overrides) == 31
    assert _salary_count_divisor(2026, 10, "fixed_days", 26, [5, 6], overrides) == 26
    assert _salary_count_divisor(2026, 10, "working_days", 26, [5, 6], overrides) == 22


def test_unconfigured_employee_does_not_get_configured_salary():
    result = _calculate_employee_monthly_payroll(
        _employee(), None, [], 2026, 10, "fixed_days", 26, [5, 6], {}
    )
    assert result["configured"] is False
    assert result["annual_salary"] == 0
    assert result["net_pay"] == 0


def test_week_off_days_validate_and_normalize():
    assert PayrollSettingsIn(week_off_days=[6, 5]).week_off_days == [5, 6]
    with pytest.raises(ValidationError):
        PayrollSettingsIn(week_off_days=[7])
    with pytest.raises(ValidationError):
        PayrollSettingsIn(week_off_days=[6, 6])


def test_absences_on_holidays_and_week_offs_do_not_create_salary_deductions():
    from app.api.routes.hr_payroll import _calculate_employee_monthly_payroll

    attendance = [
        SimpleNamespace(status="absent", work_date=date(2026, 10, 3), overtime_minutes=0),
        SimpleNamespace(status="absent", work_date=date(2026, 10, 5), overtime_minutes=0),
        SimpleNamespace(status="absent", work_date=date(2026, 10, 6), overtime_minutes=0),
    ]
    result = _calculate_employee_monthly_payroll(
        _employee(),
        _profile(),
        attendance,
        2026,
        10,
        "fixed_days",
        26,
        [5, 6],
        {date(2026, 10, 5): "holiday"},
    )

    # The three absence records remain visible, but only the Tuesday workday
    # is deductible; October 3 is a default week-off and October 5 is a holiday.
    assert result["absent_days"] == 3
    assert result["absence_deduction"] == Decimal("1923.08")


def test_absence_deduction_cannot_exceed_monthly_gross():
    attendance = [
        SimpleNamespace(
            status="absent",
            work_date=date(2026, 10, day),
            overtime_minutes=0,
        )
        for day in range(1, 32)
    ]
    result = _calculate_employee_monthly_payroll(
        _employee(), _profile(), attendance, 2026, 10, "fixed_days", 1, [5, 6], {}
    )

    assert result["absence_deduction"] == result["monthly_gross"] == 50000
    assert result["net_pay"] == 0


@pytest.mark.parametrize(("year", "month"), [(1999, 12), (2101, 1), (2026, 0), (2026, 13)])
def test_payroll_period_validation_rejects_out_of_range_values(year, month):
    from fastapi import HTTPException

    from app.api.routes.hr_payroll import _validate_period

    with pytest.raises(HTTPException) as exc:
        _validate_period(year, month)

    assert exc.value.status_code == 422


def _attendance_workbook(rows, headers=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Attendance"
    sheet.append(headers or [
        "Employee Code", "Work Date", "Status", "Overtime Minutes", "Punch In", "Punch Out", "Note",
    ])
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_attendance_excel_template_has_exact_import_sheet_and_separate_instructions():
    content = _attendance_template_workbook(2026, 10).getvalue()
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == ["Attendance", "Instructions"]
        assert list(next(workbook["Attendance"].iter_rows(min_row=1, max_row=1, values_only=True))) == [
            "Employee Code", "Work Date", "Status", "Overtime Minutes", "Punch In", "Punch Out", "Note",
        ]
        # The example is explanatory only and does not become an import row.
        assert workbook["Attendance"].max_row == 1
        assert workbook["Instructions"].max_row >= 17
    finally:
        workbook.close()


def test_attendance_excel_parser_reads_valid_rows_and_excel_times():
    workbook = _attendance_workbook([
        [" emp-001 ", date(2026, 10, 5), "Present", 30, time_type(8, 30), time_type(17, 0), "Shift A"],
        ["EMP-002", "2026-10-06", "half day", 0, None, None, ""],
    ])
    employees = {
        "EMP-001": SimpleNamespace(id="emp-1", emp_code="EMP-001", active=True),
        "EMP-002": SimpleNamespace(id="emp-2", emp_code="EMP-002", active=True),
    }
    rows = _parse_attendance_workbook(workbook, 2026, 10, employees)

    assert len(rows) == 2
    assert rows[0]["employee_id"] == "emp-1"
    assert rows[0]["status"] == "present"
    assert rows[0]["overtime_minutes"] == 30
    assert rows[0]["punch_in"].date() == date(2026, 10, 5)
    assert rows[0]["punch_in"].hour == 8
    assert rows[0]["punch_out"].hour == 17
    assert rows[1]["status"] == "half_day"
    assert rows[1]["work_date"] == date(2026, 10, 6)


def test_attendance_excel_parser_rejects_entire_file_on_unknown_employee_or_bad_period():
    workbook = _attendance_workbook([
        ["EMP-001", date(2026, 10, 5), "present", 0, None, None, ""],
        ["UNKNOWN", date(2026, 10, 6), "present", 0, None, None, ""],
    ])
    employees = {"EMP-001": SimpleNamespace(id="emp-1", emp_code="EMP-001", active=True)}
    with pytest.raises(ValueError, match="not found or is inactive"):
        _parse_attendance_workbook(workbook, 2026, 10, employees)

    wrong_month = _attendance_workbook([
        ["EMP-001", date(2026, 9, 30), "present", 0, None, None, ""],
    ])
    with pytest.raises(ValueError, match="must be in 2026-10"):
        _parse_attendance_workbook(wrong_month, 2026, 10, employees)


def test_attendance_excel_parser_rejects_duplicate_employee_day_before_writing():
    workbook = _attendance_workbook([
        ["EMP-001", date(2026, 10, 5), "present", 0, None, None, ""],
        ["EMP-001", date(2026, 10, 5), "absent", 0, None, None, ""],
    ])
    employees = {"EMP-001": SimpleNamespace(id="emp-1", emp_code="EMP-001", active=True)}
    with pytest.raises(ValueError, match="duplicate Employee Code and Work Date"):
        _parse_attendance_workbook(workbook, 2026, 10, employees)


def test_attendance_excel_parser_rejects_header_changes():
    workbook = _attendance_workbook(
        [["EMP-001", date(2026, 10, 5), "present", 0, None, None, ""]],
        headers=["Employee", "Work Date", "Status", "Overtime Minutes", "Punch In", "Punch Out", "Note"],
    )
    with pytest.raises(ValueError, match="headers must match the template exactly"):
        _parse_attendance_workbook(
            workbook, 2026, 10,
            {"EMP-001": SimpleNamespace(id="emp-1", emp_code="EMP-001", active=True)},
        )

